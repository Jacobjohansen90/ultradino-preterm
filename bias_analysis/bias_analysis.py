#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Created on Fri Sep 25 11:54:23 2026

@author: jacob
"""

import polars as pl
from omegaconf import OmegaConf, ListConfig
import numpy as np
from pathlib import Path
from tqdm import tqdm
from sklearn.metrics import roc_auc_score, recall_score, confusion_matrix
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
   

def generate_insight(model_path, bias_cfg_path):
    #Setup
    bias_cfg = OmegaConf.load(bias_cfg_path)    
    
    #Load test data
    files = sorted((Path(model_path) / "misc").glob("*.parquet"))
    test_dfs = [pl.read_parquet(file) for file in files]
    test_df = pl.concat(test_dfs, how="vertical")
    test_df = test_df.with_columns(pl.len().over("CPR_CHILD").alias("num_images"))
    #Merge deltas and convert to μm
    test_df = test_df.with_columns(((pl.col("physical_delta_x").list.mean() +
                                     pl.col("physical_delta_y").list.mean()) / 2 * 10000).alias("physical_delta"))
                                   
    #Make df per child, not per image
    test_df = test_df.unique(subset=["CPR_CHILD"])
    
    #Run bias analysis
    #bias_analysis(test_df, model_path, bias_cfg)
    
    #Run sub group analysis
    subgroup_analysis(test_df, model_path, bias_cfg)

    #Generate histograms
    

def bias_analysis(variable_df, model_path, bias_cfg):    
    bias_wb = Workbook()
    bias_wb.remove(bias_wb.active)
    
    significance_pairs = [("P (sens)", "Sensitivity"), ("P (spec)", "Specificity")]
    
    for path in sorted((Path(model_path) / 'predictions/').iterdir()):
        if path.name.endswith(".parquet"):
            GA = path.stem.split("_")[-1]
            predictions = pl.read_parquet(path)
            
            df = predictions.join(variable_df, how='left', on='CPR_CHILD')
            results = {}
            orders = {}
            for var in tqdm(bias_cfg.bias_variables, desc=f"Processing Variables, GA {GA}"):
                if isinstance(var.variable, (list, ListConfig)):
                    df = df.with_columns(pl.coalesce([pl.col(var.variable[0]), pl.col(var.variable[1])])
                                         .alias(var.variable[0]))
                    var.variable = var.variable[0]

                results[var.name] = {}
                
                for pop, pop_df in [("All Births", df), ("Non-treated", df.filter(~pl.col("treatment")))]:
                                        
                    results[var.name][pop] = {}
                    
                    var_df = pop_df.select(['CPR_CHILD', 'preds', 'preds_binary', 'preds_normalized',
                                            'label', 'threshold', var.variable])
                    var_df = var_df.filter(pl.col(var.variable).is_not_null())

                    var_df, order = categorize_variable(var_df, var)
                    orders[var.name] = ['Overall'] + order

                    report, overall_metrics = get_bias_metrics(var_df)     

                    results[var.name][pop]['Overall'] = report
                    
                    for subgroup_df in var_df.partition_by(var.variable, maintain_order=True):
                        subgroup = subgroup_df[var.variable][0]
                        report, _ = get_bias_metrics(subgroup_df, overall_metrics)
                        p_sens = permutation_test(subgroup_df, pop_df, recall_score, preds_column='preds_binary')
                        p_spec = permutation_test(subgroup_df, pop_df, specificity, preds_column='preds_binary')
                        report['P (sens)'] = p_sens
                        report['P (spec)'] = p_spec
                        results[var.name][pop][subgroup] = report

            bias_wb = report_to_excel(bias_wb, results, orders, f"GA {GA}", 'Category', significance_pairs=significance_pairs)
    bias_wb.save(Path(model_path) / "bias_analysis.xlsx")    

def get_bias_metrics(df, overall_metrics=None):
    n = df.height
    n_preterm = df['label'].sum()
    
    labels = df['label'].to_numpy()
    preds = df['preds_normalized'].to_numpy()
    preds_binary = df['preds_binary'].to_numpy()
    
    AUC = calculate_metrics(roc_auc_score, labels, preds)['value']
    sens = calculate_metrics(recall_score, labels, preds_binary)['value']
    spec = calculate_metrics(specificity, labels, preds_binary)['value']


    report = {'N': f"{n}",
               'Preterm': f"{n_preterm}",
               'AUC': f"{AUC:.3f}",
               'Sensitivity': f"{100*sens:.1f}%",
               'Missed / 100k': '-',
               'P (sens)': '-',
               'Specificity': f"{100*spec:.1f}%",
               'Extra FP / 100k': '-',
               'P (spec)': '-'}

    if overall_metrics is None:
    
        metrics = {'n': n,
                   'preterm': n_preterm,
                   'auc': AUC,
                   'sens': sens,
                   'spec': spec}

        return report, metrics

    else:
        # Differences relative to Overall
        report['N'] = f"{n} ({100 * n / overall_metrics['n']:.1f}%)"
        report['Preterm'] = f"{n_preterm} ({100 * n_preterm / overall_metrics['n_preterm']:.1f}%)"
        report['AUC'] += f" ({AUC - overall_metrics['auc']:+.3f})"
        report['Sensitivity'] += f" ({100 * (sens - overall_metrics['sens']):+.1f}%)"
        report['Specificity'] += f" ({100 * (spec - overall_metrics['spec']):+.1f}%)"
    
        subgroup_fraction = n / overall_metrics['n']
        preterm_prevalence = n_preterm / n if n else 0

        missed = 100_000 * subgroup_fraction * preterm_prevalence * (overall_metrics['sens'] - sens)

        extra_fp = 100_000 * subgroup_fraction * (1 - preterm_prevalence) * (overall_metrics['spec'] - spec)

        report['Missed / 100k'] = f"{missed:+.1f}"
        report['Extra FP / 100k'] = f"{extra_fp:+.1f}"
        
        return report, None

def calculate_metrics(metric, labels, preds, with_CI=False, n_bootstrap=10000, n_digits=3):
    
    results = {"value": round(metric(labels, preds), n_digits)}
    
    if with_CI:
        rng = np.random.default_rng()
        n = len(labels)
        bootstrap_scores = []

        for _ in range(n_bootstrap):
            idx = rng.choice(n, size=n, replace=True)

            boot_labels = labels[idx]
            boot_preds = preds[idx]

            # Skip samples where the metric cannot be calculated
            try:
                score = metric(boot_labels, boot_preds)
                if np.isfinite(score):
                    bootstrap_scores.append(score)
            except (ValueError, ZeroDivisionError):
                continue

        if bootstrap_scores:
            ci_low, ci_high = np.percentile(bootstrap_scores, [2.5, 97.5])
        else:
            ci_low, ci_high = np.nan, np.nan

        results["ci_low"] = round(ci_low, n_digits)
        results["ci_high"] = round(ci_high, n_digits)

    return results
            
def specificity(y_true, y_pred):
    tn, fp, fn, tp = confusion_matrix(y_true, y_pred, labels=[0, 1]).ravel()
    return tn / (tn + fp)           


def categorize_variable(df, var):
    if var.type == "interval":
        categories = list(var.categories)

        order = ([f"<{categories[0]}"]
                 + [f"{lower}-{upper}" for lower, upper in zip(categories[:-1], categories[1:])]
                 + [f"{categories[-1]}+"])

        expr = pl.when(pl.col(var.variable) < categories[0])
        expr = expr.then(pl.lit(order[0]))

        for lower, upper, label in zip(categories[:-1], categories[1:], order[1:]):
            expr = expr.when(pl.col(var.variable) < upper).then(pl.lit(label))

        expr = expr.otherwise(pl.lit(order[-1]))

        return df.with_columns(expr.alias(var.variable)), order
    
    elif var.type == "binary":
        categories = list(var.categories)
    
        if len(categories) != 2:
            raise ValueError(f"Binary variable '{var.variable}' must have exactly 2 categories")
    
        order = categories
    
        expr = (pl.when(pl.col(var.variable) == True)
                .then(pl.lit(categories[0]))
                .when(pl.col(var.variable) == False)
                .then(pl.lit(categories[1]))
                .otherwise(None))
    
        return df.with_columns(expr.alias(var.variable)), order
    
    elif var.type == 'centiles':
        categories = list(var.categories)

        order = ([f"0th-{categories[0]}th"] 
                 + [f"{lower}th-{upper}th" for lower, upper in zip(categories[:-1], categories[1:])]
                 + [f"{categories[-1]}th-100th"])
       
        percentile_values = df.select([pl.col(var.variable).quantile(p / 100).alias(f"p{p}") for p in categories]).row(0) 
        expr = pl.when(pl.col(var.variable) < percentile_values[0]).then(pl.lit(order[0]))
        
        for value, label in zip(percentile_values[1:], order[1:]):
            expr = expr.when(pl.col(var.variable) < value).then(pl.lit(label))

        expr = expr.otherwise(pl.lit(order[-1]))

        return df.with_columns(expr.alias(var.variable)), order
    
    elif var.type == "year_interval":
        categories = list(var.categories)
    
        order = ([f"<{categories[0]}"]
                 + [f"{lower}-{upper}" for lower, upper in zip(categories[:-1], categories[1:])]
                 + [f"{categories[-1]}+"])
    
        year = pl.col(var.variable).dt.year()
    
        expr = pl.when(year < categories[0])
        expr = expr.then(pl.lit(order[0]))
    
        for lower, upper, label in zip(categories[:-1], categories[1:], order[1:]):
            expr = expr.when(year < upper).then(pl.lit(label))
    
        expr = expr.otherwise(pl.lit(order[-1]))
    
        return df.with_columns(expr.alias(var.variable)), order
    
    else:
        raise Exception(f"Type {var.type} not implemented")


def permutation_test(subgroup_df, pop_df, metric, preds_column='preds', n_permutations=10000):
    rng = np.random.default_rng()
    
    overall_labels = pop_df['label'].to_numpy()
    overall_preds = pop_df[preds_column].to_numpy()
    
    overall_metric = metric(overall_labels, overall_preds)
    subgroup_metric = metric(subgroup_df['label'], subgroup_df[preds_column])
    n_subgroup = subgroup_df.height
    
    observed_diff = overall_metric - subgroup_metric
    
    diffs = np.empty(n_permutations)
    
    for i in range(n_permutations):
        indices = rng.choice(pop_df.height, size=n_subgroup, replace=False)
        permuted_metric = metric(overall_labels[indices], overall_preds[indices])
        diffs[i] = overall_metric - permuted_metric
    
    p_value = (1 + np.sum(np.abs(diffs) >= abs(observed_diff))) / (n_permutations + 1)
    
    return p_value
    

def subgroup_analysis(variable_df, model_path, variable_cfg):
        wb = Workbook()
        wb.remove(wb.active)
        results = {}
        for path in sorted((Path(model_path) / 'predictions/').iterdir()):
            if path.name.endswith(".parquet"):
                GA = path.stem.split("_")[-1]

                inclusion = ((pl.col("GA") // 7 >= int(GA))
                             | (~pl.col("induced") & ~pl.col("c-section"))
                             | pl.col("pprom")
                             | (pl.col("c-section") & (pl.col("c-section_during_birth") 
                                                       | pl.col("contractions_with_preterm_birth")) & ~pl.col("induced")))
                
                predictions = pl.read_parquet(path)
                df_all = predictions.join(variable_df, how='left', on='CPR_CHILD')
                df_all = (df_all.filter(inclusion).with_columns((pl.col("GA") // 7 < int(GA)).alias("preterm")))
                df_CL = df_all.filter(pl.col("CL_manual").is_not_null())
                
                results['Model'] = {}
                results['CL'] = {}
                results['Model (CL Avaliable)'] = {}
                
                for pop, pop_df in [('All patients', df_all), ('Non-treated', df_all.filter(~pl.col("treatment")))]:
                    report = get_subgroup_metrics(pop_df, variable_cfg.subgroups, CL=None)
                    results['Model'][pop] = report
                                    
                for pop, pop_df in [('All patients', df_CL), ('Non-treated', df_CL.filter(~pl.col("treatment")))]:
                    report = get_subgroup_metrics(pop_df, variable_cfg.subgroups, CL="CL_manual")
                    results['CL'][pop] = report
                    
                    report = get_subgroup_metrics(pop_df, variable_cfg.subgroups, CL=None)
                    results['Model (CL Avaliable)'][pop] = report
                    
                orders = {model: ["All"] + [subgroup.name for subgroup in variable_cfg.subgroups] for model in results}
                wb = report_to_excel(wb , results, orders, f"GA {GA}", 'Subgroup')
                wb.save(Path(model_path) / "subgroup_analysis.xlsx")    
                

def get_subgroup_metrics(df, variables, CL=None, n_bootstrap=10000):
    results = {}
    subgroups = [{"name": "All", "var": None, "value": None},
                 *[dict(var) for var in variables]]
    
    for subgroup in subgroups:
        name = subgroup["name"]
        var = subgroup["var"]
        value = subgroup["value"]
    
        if var is None:
            subgroup_df = df
        else:
            subgroup_df = df.filter((pl.col("label") == 0) | (pl.col(var) == value))
    
        labels = subgroup_df["label"].to_numpy().astype(int)
        
        if CL is None:
            preds = subgroup_df["preds_normalized"].to_numpy()
            preds_binary = subgroup_df["preds_binary"].to_numpy()
        else:
            preds = -subgroup_df[CL].to_numpy()
            preds_binary = (subgroup_df[CL] < 25).to_numpy()
        
        AUC = calculate_metrics(roc_auc_score, labels, preds, with_CI=True)
        Sens = calculate_metrics(recall_score, labels, preds_binary, with_CI=True)
        Spec = calculate_metrics(specificity, labels, preds_binary, with_CI=True)
        
        results[name] = {'AUC': f"{AUC['value']} ({AUC['ci_low']}-{AUC['ci_high']})",
                         'Sensitivity': f"{Sens['value']:.1%} ({Sens['ci_low']:.1%}-{Sens['ci_high']:.1%})",
                         'Specificity': f"{Spec['value']:.1%} ({Spec['ci_low']:.1%}-{Spec['ci_high']:.1%})",
                         'N': f"{subgroup_df.height}",
                         'N Preterm': f"{labels.sum()}"}

    return results
    
   
def report_to_excel(wb, results, orders, title, subgroup_name, 
                    significance_level=0.05, significance_pairs=None):
    ws = wb.create_sheet(title=title)

    section_font = Font(size=14, bold=True, color="FFFFFF")
    header_font = Font(bold=True)

    section_fill_gray = PatternFill("solid", fgColor="595959")
    header_fill_gray = PatternFill("solid", fgColor="B7B7B7")
    results_fill_gray = PatternFill("solid", fgColor="EDEDED")

    section_fill_blue = PatternFill("solid", fgColor="2F5597")
    header_fill_blue = PatternFill("solid", fgColor="8FAADC")
    results_fill_blue = PatternFill("solid", fgColor="EAF2F8")

    red_font = Font(color="9C0006", bold=True)
    green_font = Font(color="006100", bold=True)

    thin_border_gray = Side(style="thin", color="D9D9D9")
    thick_border_gray = Side(style="medium", color="595959")
    thin_border_blue = Side(style="thin", color="B4C7E7")
    thick_border_blue = Side(style="medium", color="2F5597")

    row = 1

    for variable_index, (variable, population_data) in enumerate(results.items()):
        populations = list(population_data.keys())
        entries = orders[variable]

        headers = [subgroup_name] + list(next(iter(next(iter(population_data.values())).values())).keys())

        columns = len(headers)
        num_populations = len(populations)
        total_columns = columns * num_populations

        if variable_index % 2 == 0:
            section_fill = section_fill_gray
            header_fill = header_fill_gray
            results_fill = results_fill_gray
            thin_side = thin_border_gray
            thick_side = thick_border_gray
        else:
            section_fill = section_fill_blue
            header_fill = header_fill_blue
            results_fill = results_fill_blue
            thin_side = thin_border_blue
            thick_side = thick_border_blue

        # Variable section heading
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)

        cell = ws.cell(row=row, column=1, value=variable)
        cell.font = section_font
        cell.fill = section_fill
        cell.alignment = Alignment(horizontal="center")

        row += 1
        section_start_row = row

        # Population headings
        for i, population in enumerate(populations):
            start_col = i * columns + 1
            end_col = start_col + columns - 1

            ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)

            cell = ws.cell(row=row, column=start_col, value=population)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        row += 1

        # Column headers for each population
        for i in range(num_populations):
            start_col = i * columns + 1

            for j, header in enumerate(headers):
                cell = ws.cell(row=row, column=start_col + j, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

        row += 1

        # Bonferroni correction across categories, excluding Overall
        reference_entry = entries[0]
        num_categories = len(entries) - 1
        significance_threshold = (significance_level / num_categories if num_categories > 0 else significance_level)

        # Write category results
        for entry in entries:
            for i, population in enumerate(populations):
                start_col = i * columns + 1
                data = population_data[population][entry]
        
                # Category name
                cell = ws.cell(row=row, column=start_col, value=entry)
                cell.fill = results_fill
        
                # Metric values
                for j, header in enumerate(headers[1:], start=1):
                    cell = ws.cell(
                        row=row,
                        column=start_col + j,
                        value=data[header]
                    )
                    cell.fill = results_fill
        
                # Color significant metric differences vs. reference category
                if entry != reference_entry:
                    reference = population_data[population][reference_entry]
                    if significance_pairs is not None:
                        for p_key, metric in significance_pairs:
                            if p_key not in data or metric not in data:
                                continue
            
                            p_value = data[p_key]
            
                            try:
                                p_value = float(p_value)
                            except (ValueError, TypeError):
                                continue
            
                            if p_value >= significance_threshold:
                                continue
            
                            category_value = data[metric]
                            reference_value = reference[metric]
            
                            try:
                                category_value = float(
                                    str(category_value).split()[0].rstrip("%")
                                )
                                reference_value = float(
                                    str(reference_value).split()[0].rstrip("%")
                                )
                            except (ValueError, TypeError):
                                continue
            
                            if category_value < reference_value:
                                font = red_font
                            elif category_value > reference_value:
                                font = green_font
                            else:
                                continue
            
                            p_col = start_col + headers.index(p_key)
                            ws.cell(row=row, column=p_col).font = font
        
            row += 1

        section_end_row = row - 1

        # Borders around each population's table
        for i in range(num_populations):
            start_col = i * columns + 1
            end_col = start_col + columns - 1

            for r in range(section_start_row, section_end_row + 1):
                for c in range(start_col, end_col + 1):
                    cell = ws.cell(row=r, column=c)

                    cell.border = Border(
                        left=thick_side if c == start_col else thin_side,
                        right=thick_side if c == end_col else thin_side,
                        top=thick_side if r == section_start_row else thin_side,
                        bottom=thick_side if r == section_end_row else thin_side
                    )

    # Adjust column widths
    for column in ws.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )

        column_letter = get_column_letter(column[0].column)
        ws.column_dimensions[column_letter].width = max_length + 2

    return wb
    