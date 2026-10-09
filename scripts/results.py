import argparse

import numpy as np
import polars as pl

from tqdm import tqdm
from sklearn.metrics import roc_auc_score, roc_curve

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

parser = argparse.ArgumentParser()

parser.add_argument(
    "--model_name",
    type=str,
    required=True,
    help="Name of the model/experiment",
)

args = parser.parse_args()

model = args.model_name

bias_df = pl.read_csv('/users/data/UCPH/DeepFetal/projects/preterm/Data/misc/bias_variables.csv')

cutoff = model.split('_')[-1]

output_path = f"/users/data/UCPH/DeepFetal/projects/preterm/misc/{cutoff}.xlsx"

def sens_at_85_spec(y_true, y_score, target_spec=0.85):
    fpr, tpr, thresholds = roc_curve(y_true, y_score)

    specificity = 1 - fpr

    valid = specificity >= target_spec

    if not np.any(valid):
        return np.nan

    return np.max(tpr[valid])


def bootstrap_metrics(
    df,
    score_col,
    lower_is_positive=False,
    n_bootstrap=2000,
    seed=42,
):
    data = df.select(["label", score_col]).drop_nulls()

    y = data["label"].to_numpy()
    scores = data[score_col].to_numpy()

    if lower_is_positive:
        scores = -scores

    # Point estimates
    auc = roc_auc_score(y, scores)
    sens = sens_at_85_spec(y, scores)

    # Bootstrap
    rng = np.random.default_rng(seed)

    auc_boot = []
    sens_boot = []

    n = len(data)

    for _ in range(n_bootstrap):

        idx = rng.integers(0, n, n)

        y_b = y[idx]
        scores_b = scores[idx]

        if len(np.unique(y_b)) < 2:
            continue

        auc_boot.append(
            roc_auc_score(y_b, scores_b)
        )

        sens_boot.append(
            sens_at_85_spec(y_b, scores_b)
        )

    auc_ci = np.percentile(auc_boot, [2.5, 97.5])
    sens_ci = np.percentile(sens_boot, [2.5, 97.5])

    return {
        "n": n,
        "n_positive": int(y.sum()),
        "n_negative": int((y == 0).sum()),
        "auc": auc,
        "auc_ci_lower": auc_ci[0],
        "auc_ci_upper": auc_ci[1],
        "sens_at_85_spec": sens,
        "sens_ci_lower": sens_ci[0],
        "sens_ci_upper": sens_ci[1],
    }

# ---------------------------------------------------------------------
# 1. Collapse df to one row per child and join with predictions
# ---------------------------------------------------------------------

cutoff = int(model.split('_')[-1])
path = '/users/data/UCPH/DeepFetal/projects/preterm/training_runs/Running/'
preds = pl.read_parquet(path + model + f"/results/predictions/predictions_{cutoff}.parquet")
df = pl.read_parquet('/users/data/UCPH/DeepFetal/projects/preterm/Data/dataset_v6/test2.parquet')

df_combined = (
    preds
    .join(
        df.group_by("CPR_CHILD").agg([
            pl.col("GA").first(),
            pl.col("induced").first(),
            pl.col("c-section").first(),
            pl.col("pprom").first(),
            pl.col("c-section_during_birth").first(),
            pl.col("contractions_with_preterm_birth").first(),
            pl.col("CL").first(),
        ]),
        on="CPR_CHILD",
        how="left",
    )
    .filter(
        (pl.col("GA") // 7 >= cutoff)
        | (~pl.col("induced") & ~pl.col("c-section"))
        | pl.col("pprom")
        | (
            pl.col("c-section")
            & (
                pl.col("c-section_during_birth")
                | pl.col("contractions_with_preterm_birth")
            )
            & ~pl.col("induced")
        )
    )
)

df_combined = (
    df_combined
    .join(
        bias_df.select([
            pl.col("b_cpr").alias("CPR_CHILD"),
            "maternal_BMI",
            "maternal_age",
            "fertility_treatment_2_years_prior",
            "smoking_status"
        ]),
        on="CPR_CHILD",
        how="left",
    )
)

populations = {
    "All": pl.lit(True),
    "Non-treated": ~pl.col("treatment"),
}

sub_groups = {
    "All": None,
    "Preterm": None,
    "Non-preterm": None,
    "PPROM": "pprom",
    "C-Section": "c-section_during_birth",
}


# ---------------------------------------------------------------------
# 4. Calculate metrics
# ---------------------------------------------------------------------

results = {}
bias_results = {}

total = len(populations) * len(sub_groups)

with tqdm(total=total, desc="Calculating metrics") as pbar:

    for population_name, population_filter in populations.items():

        results[population_name] = {}
        bias_results[population_name] = {}

        population_df = df_combined.filter(population_filter)

        for sub_group_name, sub_group_value in sub_groups.items():

            # =============================================================
            # Define subgroup
            # =============================================================

            if sub_group_name == "All":

                subgroup_df = population_df
                demographic_df = subgroup_df


            elif sub_group_name == "Preterm":

                subgroup_df = population_df.filter(
                    pl.col("GA") // 7 < cutoff
                )
                demographic_df = subgroup_df

            elif sub_group_name == "Non-preterm":

                subgroup_df = population_df.filter(
                    pl.col("GA") // 7 >= cutoff
                )
                demographic_df = subgroup_df

            else:

                subgroup_df = population_df.filter(
                    (pl.col("GA") // 7 >= cutoff)
                    | (
                        (pl.col("GA") // 7 < cutoff)
                        & pl.col(sub_group_value)
                    )
                )
                demographic_df = population_df.filter((pl.col("GA") // 7 < cutoff)
                        & pl.col(sub_group_value)
                    )



            # =============================================================
            # Demographics
            # =============================================================

            total_patients = len(demographic_df)

            # -------------------------------------------------------------
            # BMI
            # -------------------------------------------------------------

            bmi = demographic_df["maternal_BMI"].drop_nulls()

            # -------------------------------------------------------------
            # Age
            # -------------------------------------------------------------

            age = demographic_df["maternal_age"].drop_nulls()

            # -------------------------------------------------------------
            # Fertility treatment
            # -------------------------------------------------------------

            fertility = demographic_df[
                "fertility_treatment_2_years_prior"
            ].drop_nulls()

            fertility_n = int(fertility.sum())

            # -------------------------------------------------------------
            # Smoking
            # -------------------------------------------------------------

            smoking = (
                demographic_df
                .select(
                    pl.col("smoking_status")
                    .cast(pl.String)
                    .alias("smoking_status")
                )
                .with_columns(
                    pl.when(
                        (pl.col("smoking_status") == "-1")
                        | pl.col("smoking_status").str.ends_with("99")
                    )
                    .then(None)
                    .when(
                        pl.col("smoking_status").str.ends_with("00")
                    )
                    .then(False)
                    .otherwise(True)
                    .alias("smoking_binary")
                )
                .filter(
                    pl.col("smoking_binary").is_not_null()
                )
            )

            smoking_n = len(smoking)
            smoking_true = smoking["smoking_binary"].sum()

            # -------------------------------------------------------------
            # Store demographics
            # -------------------------------------------------------------

            bias_results[population_name][sub_group_name] = {
                "Total patients": total_patients,

                "BMI": {
                    "non_null": len(bmi),
                    "value": (
                        f"{bmi.mean():.1f} ({bmi.std():.1f})"
                        if len(bmi) > 0
                        else "-"
                    ),
                },

                "Age": {
                    "non_null": len(age),
                    "value": (
                        f"{age.mean():.1f} ({age.std():.1f})"
                        if len(age) > 0
                        else "-"
                    ),
                },

                "Fertility treatment": {
                    "non_null": len(fertility),
                    "value": (
                        f"{fertility_n} "
                        f"({100 * fertility_n / total_patients:.1f}%)"
                        if total_patients > 0
                        else "-"
                    ),
                },

                "Smoking": {
                    "non_null": smoking_n,
                    "value": (
                        f"{smoking_true} "
                        f"({100 * smoking_true / total_patients:.1f}%)"
                        if total_patients > 0
                        else "-"
                    ),
                },
            }

            # =============================================================
            # Model metrics
            #
            # Only calculate these for the original model subgroups:
            # All, PPROM and C-Section
            # =============================================================

            if sub_group_name in ["All", "PPROM", "C-Section"]:

                # ---------------------------------------------------------
                # Model — all patients
                # ---------------------------------------------------------

                model_results = bootstrap_metrics(
                    subgroup_df,
                    score_col="preds",
                    lower_is_positive=False,
                )

                # ---------------------------------------------------------
                # Patients with a valid CL measurement
                # ---------------------------------------------------------

                cl_df = subgroup_df.filter(
                    pl.col("CL") != 0
                )

                # ---------------------------------------------------------
                # Model — CL available
                # ---------------------------------------------------------

                model_cl_results = bootstrap_metrics(
                    cl_df,
                    score_col="preds",
                    lower_is_positive=False,
                )

                # ---------------------------------------------------------
                # CL — CL available
                # ---------------------------------------------------------

                cl_results = bootstrap_metrics(
                    cl_df,
                    score_col="CL",
                    lower_is_positive=True,
                )

                results[population_name][sub_group_name] = {
                    "Model": model_results,
                    "Model (CL available)": model_cl_results,
                    "CL": cl_results,
                }

            pbar.update(1)

wb = Workbook()

ws = wb.active
ws.title = "Metrics"


# =====================================================================
# 9. Excel formatting
# =====================================================================

thin_gray = Side(
    style="thin",
    color="D9D9D9",
)

medium_gray = Side(
    style="medium",
    color="A6A6A6",
)

header_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

section_fill = PatternFill(
    fill_type="solid",
    fgColor="E7E6E6",
)

subsection_fill = PatternFill(
    fill_type="solid",
    fgColor="F2F2F2",
)

title_font = Font(
    bold=True,
    size=14,
)

section_font = Font(
    bold=True,
    size=11,
)

header_font = Font(
    bold=True,
)

normal_alignment = Alignment(
    horizontal="center",
    vertical="center",
)

left_alignment = Alignment(
    horizontal="left",
    vertical="center",
)

center_alignment = Alignment(
    horizontal="center",
    vertical="center",
)

# =====================================================================
# 10. Excel helper functions
# =====================================================================

def metric_string(
    result,
    value_key,
    lower_key,
    upper_key,
):
    return (
        f"{result[value_key]:.3f} "
        f"({result[lower_key]:.3f}-{result[upper_key]:.3f})"
    )


def write_metric_row(
    ws,
    row,
    start_col,
    subgroup_name,
    result,
):

    ws.cell(
        row=row,
        column=start_col,
        value=subgroup_name,
    )

    ws.cell(
        row=row,
        column=start_col + 1,
        value=metric_string(
            result,
            "auc",
            "auc_ci_lower",
            "auc_ci_upper",
        ),
    )

    ws.cell(
        row=row,
        column=start_col + 2,
        value=metric_string(
            result,
            "sens_at_85_spec",
            "sens_ci_lower",
            "sens_ci_upper",
        ),
    )

    ws.cell(
        row=row,
        column=start_col + 3,
        value=result["n_positive"],
    )

    ws.cell(
        row=row,
        column=start_col + 4,
        value=result["n"],
    )


def format_metric_block(
    ws,
    title_row,
    population_row,
    header_row,
    data_start_row,
    population_left,
    population_right,
):

    # Section title
    ws.merge_cells(
        start_row=title_row,
        start_column=1,
        end_row=title_row,
        end_column=10,
    )

    ws.cell(
        title_row,
        1,
    ).font = section_font

    ws.cell(
        title_row,
        1,
    ).fill = section_fill

    ws.cell(
        title_row,
        1,
    ).alignment = left_alignment

    # Population headers
    ws.merge_cells(
        start_row=population_row,
        start_column=1,
        end_row=population_row,
        end_column=5,
    )

    ws.merge_cells(
        start_row=population_row,
        start_column=6,
        end_row=population_row,
        end_column=10,
    )

    ws.cell(
        population_row,
        1,
        population_left,
    )

    ws.cell(
        population_row,
        6,
        population_right,
    )

    for col in [1, 6]:

        cell = ws.cell(
            population_row,
            col,
        )

        cell.font = header_font
        cell.fill = subsection_fill
        cell.alignment = normal_alignment

    # Column headers
    headers = [
        "Subgroup",
        "AUC",
        "Sens@Spec",
        "N-preterm",
        "N-total",
    ]

    for i, header in enumerate(headers, start=1):

        ws.cell(
            header_row,
            i,
            header,
        )

        ws.cell(
            header_row,
            i,
        ).font = header_font

        ws.cell(
            header_row,
            i,
        ).fill = header_fill

        ws.cell(
            header_row,
            i,
        ).alignment = normal_alignment

        ws.cell(
            header_row,
            i + 5,
            header,
        )

        ws.cell(
            header_row,
            i + 5,
        ).font = header_font

        ws.cell(
            header_row,
            i + 5,
        ).fill = header_fill

        ws.cell(
            header_row,
            i + 5,
        ).alignment = normal_alignment

    # Data rows
    subgroup_names = [
        "All",
        "PPROM",
        "C-Section",
    ]

    for i, subgroup_name in enumerate(
        subgroup_names,
        start=data_start_row,
    ):

        for col in range(1, 11):

            cell = ws.cell(
                i,
                col,
            )

            cell.alignment = normal_alignment

            cell.border = Border(
                bottom=thin_gray,
            )


# =====================================================================
# 11. Title
# =====================================================================

ws["A1"] = f"GA {cutoff}"
ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws["A1"].fill = PatternFill(fill_type="solid", fgColor="595959")
ws["A1"].alignment = center_alignment

ws.merge_cells("A1:J1")


# =====================================================================
# 12. Model section
# =====================================================================

format_metric_block(
    ws,
    title_row=2,
    population_row=3,
    header_row=4,
    data_start_row=5,
    population_left="Model — All patients",
    population_right="Model — Non-treatment",
)

for i, subgroup_name in enumerate(
    ["All", "PPROM", "C-Section"],
    start=5,
):
    write_metric_row(
        ws, i, 1, subgroup_name,
        results["All"][subgroup_name]["Model"],
    )

    write_metric_row(
        ws, i, 6, subgroup_name,
        results["Non-treated"][subgroup_name]["Model"],
    )


# =====================================================================
# 13. Cervical Length section
# =====================================================================

format_metric_block(
    ws,
    title_row=8,
    population_row=9,
    header_row=10,
    data_start_row=11,
    population_left="CL — All patients",
    population_right="CL — Non-treatment",
)

for i, subgroup_name in enumerate(
    ["All", "PPROM", "C-Section"],
    start=11,
):
    write_metric_row(
        ws, i, 1, subgroup_name,
        results["All"][subgroup_name]["CL"],
    )

    write_metric_row(
        ws, i, 6, subgroup_name,
        results["Non-treated"][subgroup_name]["CL"],
    )


# =====================================================================
# 14. Model — CL patients section
# =====================================================================

format_metric_block(
    ws,
    title_row=14,
    population_row=15,
    header_row=16,
    data_start_row=17,
    population_left="Model (CL available) — All patients",
    population_right="Model (CL available) — Non-treatment")

for i, subgroup_name in enumerate(
    ["All", "PPROM", "C-Section"],
    start=17,
):
    write_metric_row(
        ws,
        i,
        1,
        subgroup_name,
        results["All"][subgroup_name]["Model (CL available)"],
    )

    write_metric_row(
        ws,
        i,
        6,
        subgroup_name,
        results["Non-treated"][subgroup_name]["Model (CL available)"],
    )


# =====================================================================
# 15. General formatting
# =====================================================================

widths = {
    "A": 18,
    "B": 22,
    "C": 22,
    "D": 14,
    "E": 14,
    "F": 18,
    "G": 22,
    "H": 22,
    "I": 14,
    "J": 14,
}

for column, width in widths.items():
    ws.column_dimensions[column].width = width


# Center all cells
for row in ws.iter_rows(
    min_row=1,
    max_row=ws.max_row,
    min_col=1,
    max_col=10,
):
    for cell in row:
        if cell.value is not None:
            cell.alignment = center_alignment


# Stronger borders around population blocks
for start_row, end_row in [
    (3, 7),
    (9, 13),
    (15, 19),
]:
    for row in range(start_row, end_row + 1):
        for col in [1, 5, 6, 10]:
            ws.cell(row, col).border = Border(
                left=medium_gray if col in [1, 6] else thin_gray,
                right=medium_gray if col in [5, 10] else thin_gray,
                bottom=thin_gray,
            )


# Freeze title
ws.freeze_panes = "A2"

# =====================================================================
# 16. Demographics / Bias Analysis
# =====================================================================

ws_bias = wb.create_sheet("Demographics")

demographic_gray = PatternFill(
    fill_type="solid",
    fgColor="F2F2F2",
)

demographic_blue = PatternFill(
    fill_type="solid",
    fgColor="EAF2F8",
)

# =====================================================================
# Title
# =====================================================================

ws_bias["A1"] = f"GA {cutoff}"
ws_bias["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws_bias["A1"].fill = PatternFill(
    fill_type="solid",
    fgColor="595959",
)
ws_bias["A1"].alignment = center_alignment

ws_bias.merge_cells("A1:E1")


# =====================================================================
# Population headers
# =====================================================================

ws_bias.merge_cells("B3:C3")
ws_bias.merge_cells("D3:E3")

ws_bias["B3"] = "All patients"
ws_bias["D3"] = "Non-treatment"

for col in ["B3", "D3"]:
    ws_bias[col].font = header_font
    ws_bias[col].fill = subsection_fill
    ws_bias[col].alignment = center_alignment


# =====================================================================
# Column headers
# =====================================================================

headers = [
    "Variable",
    "Value (SD) / Count (%)",
    "N patients (% of total population)",
    "Value (SD) / Count (%)",
    "N patients (% of total population)",
]

for i, header in enumerate(headers, start=1):
    cell = ws_bias.cell(
        row=4,
        column=i,
        value=header,
    )
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_alignment


# =====================================================================
# Data
# =====================================================================

subgroups = [
    "All",
    "Preterm",
    "Non-preterm",
    "PPROM",
    "C-Section",
]

variables = [
    "BMI",
    "Age",
    "Fertility treatment",
    "Smoking",
]

row = 5

for variable_idx, variable_name in enumerate(variables):

    # Alternate colour for each variable block
    fill = (
        demographic_gray
        if variable_idx % 2 == 0
        else demographic_blue
    )

    for subgroup_name in subgroups:

        result_all = bias_results["All"][subgroup_name]
        total_all = result_all["Total patients"]

        result_non_treated = bias_results["Non-treated"][subgroup_name]
        total_non_treated = result_non_treated["Total patients"]

        # -------------------------------------------------------------
        # All patients
        # -------------------------------------------------------------

        all_value = result_all[variable_name]["value"]
        all_non_null = result_all[variable_name]["non_null"]

        all_n = (
            f"{all_non_null} "
            f"({100 * all_non_null / total_all:.1f}%)"
            if total_all > 0
            else "-"
        )

        # -------------------------------------------------------------
        # Non-treatment
        # -------------------------------------------------------------

        non_treated_value = result_non_treated[variable_name]["value"]
        non_treated_non_null = result_non_treated[variable_name]["non_null"]

        non_treated_n = (
            f"{non_treated_non_null} "
            f"({100 * non_treated_non_null / total_non_treated:.1f}%)"
            if total_non_treated > 0
            else "-"
        )

        # -------------------------------------------------------------
        # Row
        # -------------------------------------------------------------

        values = [
            f"{variable_name} ({subgroup_name})",
            all_value,
            all_n,
            non_treated_value,
            non_treated_n,
        ]

        for col, value in enumerate(values, start=1):

            cell = ws_bias.cell(
                row=row,
                column=col,
                value=value,
            )

            cell.fill = fill
            cell.alignment = center_alignment
            cell.border = Border(
                bottom=thin_gray
            )

        row += 1


# =====================================================================
# Column widths
# =====================================================================

widths = {
    "A": 30,
    "B": 24,
    "C": 35,
    "D": 24,
    "E": 35,
}

for column, width in widths.items():
    ws_bias.column_dimensions[column].width = width


# =====================================================================
# Borders
# =====================================================================

for row_idx in range(3, ws_bias.max_row + 1):

    for col in range(1, 6):

        ws_bias.cell(row_idx, col).border = Border(
            left=(
                medium_gray
                if col in [1, 2, 4]
                else thin_gray
            ),
            right=(
                medium_gray
                if col in [1, 3, 5]
                else thin_gray
            ),
            bottom=thin_gray,
        )


# =====================================================================
# General formatting
# =====================================================================

for row_cells in ws_bias.iter_rows(
    min_row=1,
    max_row=ws_bias.max_row,
    min_col=1,
    max_col=5,
):

    for cell in row_cells:

        if cell.value is not None:
            cell.alignment = center_alignment


# =====================================================================
# Freeze panes
# =====================================================================

ws_bias.freeze_panes = "A5"

# =====================================================================
# 17. Save
# =====================================================================

wb.save(output_path)

print(f"\nSaved Excel file to: {output_path}")